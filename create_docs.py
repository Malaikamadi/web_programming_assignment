import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = Document()
    
    # Page setup: Standard Letter/A4 with 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        # Configure footer with copyright
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("© Alimatu Maliaka Jalloh, BSC(Hons) Software Engineering with Multimedia, Semester 5, 905004463, 2026")
        f_run.font.name = "Tahoma"
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(100, 116, 139)

    # Set normal style font to Tahoma 12pt
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Tahoma'
    font.size = Pt(12)
    font.color.rgb = RGBColor(30, 41, 59)

    # =========================================================================
    # PAGE 1: COVER PAGE
    # =========================================================================
    
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("LIMKOKWING UNIVERSITY OF CREATIVE TECHNOLOGY\nSIERRA LEONE CAMPUS")
    r_inst.font.name = "Tahoma"
    r_inst.font.size = Pt(16)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(15, 23, 42)
    p_inst.paragraph_format.space_after = Pt(8)

    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_dept = p_dept.add_run("FACULTY OF INFORMATION & COMMUNICATION TECHNOLOGY\nDEPARTMENT OF SOFTWARE ENGINEERING & MULTIMEDIA")
    r_dept.font.name = "Tahoma"
    r_dept.font.size = Pt(11)
    r_dept.font.color.rgb = RGBColor(71, 85, 105)
    p_dept.paragraph_format.space_after = Pt(36)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title_badge = p_title.add_run("COURSE ASSIGNMENT 1 (WEEKS 1 TO 4)\n")
    r_title_badge.font.name = "Tahoma"
    r_title_badge.font.size = Pt(12)
    r_title_badge.font.bold = True
    r_title_badge.font.color.rgb = RGBColor(225, 29, 72)
    
    r_title = p_title.add_run("Design and Development of Web Platform & Authentication System for Django Girls Sierra Leone")
    r_title.font.name = "Tahoma"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    p_title.paragraph_format.space_after = Pt(40)

    # Decorative Box for Submission Metadata
    table_meta = doc.add_table(rows=6, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.autofit = False

    meta_rows = [
        ("Student Name:", "Alimatu Maliaka Jalloh"),
        ("Student ID:", "905004463"),
        ("Academic Program:", "BSC(Hons) Software Engineering with Multimedia"),
        ("Current Semester:", "Semester 5"),
        ("Module / Course:", "Web Programming (DTECH 206 / ICT204)"),
        ("Lecturer in Charge:", "Kenneth G. Campbell")
    ]

    for i, (label, val) in enumerate(meta_rows):
        row = table_meta.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(2.2)
        c2.width = Inches(4.3)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.name = "Tahoma"
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(30, 41, 59)
        p1.paragraph_format.space_after = Pt(4)

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.name = "Tahoma"
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(51, 65, 85)
        p2.paragraph_format.space_after = Pt(4)

        set_cell_background(c1, "FFF1F2")
        set_cell_background(c2, "FFFFFF")
        set_cell_margins(c1, top=80, bottom=80, left=100, right=100)
        set_cell_margins(c2, top=80, bottom=80, left=100, right=100)

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(40)
    r_date = p_date.add_run("Submission Academic Year: 2026\nFreetown, Sierra Leone")
    r_date.font.name = "Tahoma"
    r_date.font.size = Pt(10.5)
    r_date.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_page_break()

    # =========================================================================
    # PAGE 2: SECTION 1 & SECTION 2
    # =========================================================================
    
    def add_section_header(num, title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r_num = p.add_run(f"{num}. ")
        r_num.font.name = "Tahoma"
        r_num.font.size = Pt(14)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(225, 29, 72)
        
        r_t = p.add_run(title)
        r_t.font.name = "Tahoma"
        r_t.font.size = Pt(14)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(15, 23, 42)
        return p

    def add_body_paragraph(text, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Tahoma"
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    add_section_header("1", "Executive Summary & Selected Organisation Profile")
    add_body_paragraph(
        "For Assignment 1, the chosen organisation is Django Girls Sierra Leone (Freetown Chapter). "
        "Django Girls is an inspiring, non-profit volunteer movement dedicated to bridging the digital gender divide by introducing "
        "women and young girls to computer programming and modern software craftsmanship through completely free, high-impact workshops. "
        "In Sierra Leone, female representation in computing, software engineering, and technical leadership remains historically low. "
        "Django Girls Freetown solves this challenge by dismantling intimidating barriers to entry and providing an inclusive, safe, and supportive "
        "environment where beginners can transition from tech consumers to software creators in a single day."
    )
    add_body_paragraph(
        "The digital platform developed for this assignment provides the core web infrastructure for Django Girls Sierra Leone. "
        "It incorporates modern HTML5 structure, bespoke CSS aesthetics, server-side PHP data handling, MySQL database persistence, "
        "and secure user session management, directly reflecting the learning outcomes covered in Weeks 1 to 4 of the Web Programming syllabus."
    )

    add_section_header("2", "Architecture & Web Platform Page Structure")
    add_body_paragraph(
        "The application is structured into modular components adhering to web development industry standards:"
    )

    modules = [
        ("• Prominent Logo Display (Header): ", "The official Django Girls logo is positioned prominently at the top header of every page alongside the Sierra Leone national tech badge, establishing instantaneous brand recognition and organizational legitimacy."),
        ("• Home Page (index.php): ", "Serves as the public-facing gateway showcasing the organisation's mission, impact statistics (500+ girls trained in SL, 100% free), interactive code terminals, live upcoming workshop cohorts dynamically queried from MySQL, and calls-to-action."),
        ("• User Registration Page (register.php): ", "Provides a complete onboarding registration form capturing full name, username, email, phone number, experience level, and preferred learning track, backed by comprehensive server-side validation and password encryption."),
        ("• Secure Authentication Page (login.php): ", "Implements a secure login form utilizing PDO prepared statements to query user credentials, verifying input passwords against encrypted database hashes and establishing a protected PHP session."),
        ("• Personalized User Dashboard (dashboard.php): ", "A protected portal accessible only after successful authentication. Displays tailored profile information, account status, enrolled training progress, assigned mentors, and quick administrative shortcuts."),
        ("• Secure Session Termination (logout.php): ", "Securely unsets session memory, destroys session records on the server, removes browser session cookies, and safely redirects the user back to the login interface with flash feedback."),
        ("• Standardized Copyright Footer (footer.php): ", "Incorporates the mandatory assignment footer displaying: '© Alimatu Maliaka Jalloh, BSC(Hons) Software Engineering with Multimedia, Semester 5, 905004463, 2026' on every page.")
    ]

    for label, desc in modules:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(label)
        r1.font.name = "Tahoma"
        r1.font.size = Pt(12)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(15, 23, 42)
        r2 = p.add_run(desc)
        r2.font.name = "Tahoma"
        r2.font.size = Pt(12)
        r2.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

    # =========================================================================
    # PAGE 3: SECTION 3 & SECTION 4
    # =========================================================================

    add_section_header("3", "Detailed Explanation of Design Choices, Plugins & Themes")
    add_body_paragraph(
        "A deliberate, user-centered design system was engineered specifically for the Django Girls Sierra Leone platform. "
        "The interface prioritizes simplicity, professionalism, and a distinctive feminine aesthetic tailored to women in technology:"
    )

    design_choices = [
        ("Feminine & Professional Color Palette: ", "The visual identity is styled around a chic, modern feminine palette featuring rose-pink (#e11d48) and coral-pink (#fb7185) gradients as primary accents. This is paired with delicate rose-blush surfaces (#fff1f2, #faf8f9) and deep midnight-plum slate (#181124) for headers and cards, producing an empowering, sophisticated look that celebrates women in STEM."),
        ("Professional Vector SVG Icons (No Emojis): ", "To maintain a high standard of professional visual hierarchy, all decorative OS emojis were replaced with crisp, lightweight inline SVG vector icons (calendar, map pin, sparkle, code, users, globe, shield, lock). This guarantees uniform rendering across all mobile, desktop, and operating system platforms."),
        ("Clean, Minimalist Layout Structure: ", "Cards and containers feature balanced padding, smooth rounded borders (14px radius), soft shadows (box-shadow: 0 4px 10px rgba(225, 29, 72, 0.05)), and subtle borders (#fce7f3), avoiding visual clutter and keeping user attention focused on key actions."),
        ("Typography Strategy: ", "The platform employs Plus Jakarta Sans and Tahoma as primary font families, ensuring crisp legibility and strict assignment compliance. Accent script (Caveat) is selectively used on the logo mark to preserve community friendliness."),
        ("Interactive Feedback & HCI Principles: ", "Buttons feature smooth micro-interactions on hover, form fields provide active rose focus rings, and registration fields feature real-time password visibility toggles and matching status indicators.")
    ]

    for title, desc in design_choices:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(f"• {title}")
        r1.font.name = "Tahoma"
        r1.font.size = Pt(12)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(15, 23, 42)
        r2 = p.add_run(desc)
        r2.font.name = "Tahoma"
        r2.font.size = Pt(12)
        r2.font.color.rgb = RGBColor(51, 65, 85)

    add_section_header("4", "PHP Technical Implementation & Security Considerations")
    add_body_paragraph(
        "The server-side backend demonstrates proficiency in foundational PHP concepts covered in Weeks 1 to 4:"
    )

    security_points = [
        ("Password Hashing (password_hash & password_verify): ", "Plaintext passwords are never stored. During user registration, the system applies PHP's native password_hash($password, PASSWORD_DEFAULT), producing an irreversible 60-character bcrypt hash with automated cryptographic salting. Authentication uses password_verify() to evaluate login attempts securely in constant time."),
        ("SQL Injection Prevention via PDO Prepared Statements: ", "Database operations utilize PHP Data Objects (PDO). All dynamic user inputs are executed through prepared statements with parameter binding (e.g., $stmt->execute([$username])), eliminating SQL injection vulnerabilities."),
        ("Session Protection & Fixation Defense: ", "Upon successful credential verification, session_regenerate_id(true) is invoked to replace the active session identifier with a new cryptographically secure token, preventing session fixation and hijacking attacks. Session destruction in logout.php completely wipes $_SESSION arrays and deletes client session cookies."),
        ("Cross-Site Scripting (XSS) Sanitization: ", "All user-generated data displayed across the public pages and personalized dashboard is filtered through htmlspecialchars() with UTF-8 encoding, safeguarding against script injection attacks.")
    ]

    for title, desc in security_points:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(f"• {title}")
        r1.font.name = "Tahoma"
        r1.font.size = Pt(12)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(15, 23, 42)
        r2 = p.add_run(desc)
        r2.font.name = "Tahoma"
        r2.font.size = Pt(12)
        r2.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

    # =========================================================================
    # PAGE 4: SECTION 5 - PRINTED SCREENSHOT OF HOME PAGE ONLY
    # =========================================================================

    add_section_header("5", "Home Page Interface Presentation (Printed Screenshot)")
    add_body_paragraph(
        "In strict compliance with the Assignment 1 submission guidelines ('Include printed screenshots of the following pages: Home Page only'), "
        "the figure below presents the fully rendered Home Page interface of Django Girls Sierra Leone as served locally via XAMPP Apache."
    )

    # Insert Screenshot if file exists
    screenshot_path = r"c:\Users\madi\Desktop\Web_Programming_assignment\assets\images\homepage_screenshot.png"
    if os.path.exists(screenshot_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(6)
        run_img = p_img.add_run()
        run_img.add_picture(screenshot_path, width=Inches(6.2))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run("Figure 1.0: Printed Screenshot of the Django Girls Sierra Leone Home Page (index.php) displaying prominent logo, feminine rose palette, vector icons, and statistics.")
        r_cap.font.name = "Tahoma"
        r_cap.font.size = Pt(10)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(100, 116, 139)

    add_body_paragraph(
        "Key Home Page Interface Observations:\n"
        "1. Prominent Logo Display: The Django Girls logo is positioned in the sticky header bar with custom coral highlights and the Sierra Leone badge.\n"
        "2. Feminine & Professional Aesthetic: Warm rose accents, crisp SVG icons, and elegant card styling give the site a distinctive feminine yet professional atmosphere.\n"
        "3. Live Cohort Database Integration: Below the hero section, active workshop sessions are queried directly from the MySQL database and presented in responsive cards."
    )

    doc.add_page_break()

    # =========================================================================
    # PAGE 5: SECTION 6 - INDIVIDUAL ASSESSMENT FORM & DECLARATION
    # =========================================================================

    add_section_header("6", "Individual Assessment Form & Student Declaration")
    add_body_paragraph(
        "As required by the assignment guidelines, this individual assessment form outlines the evaluation criteria, "
        "deliverables achieved, and personal declaration of original work."
    )

    # Assessment Rubric Table
    rubric_table = doc.add_table(rows=8, cols=3)
    rubric_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rubric_table.autofit = False

    rubric_headers = ["Assessment Criteria / Deliverable", "Requirement Description", "Status & Evidence"]
    for j, h in enumerate(rubric_headers):
        cell = rubric_table.rows[0].cells[j]
        set_cell_background(cell, "181124")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Tahoma"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)

    rubric_data = [
        ("Logo Prominence", "Prominently display the Django Girls logo at the top header.", "Achieved (Visible in sticky header)"),
        ("Sign-Up Page", "User registration form with input validation and database insert.", "Achieved (register.php with validation)"),
        ("Login Page", "Authentication form verifying user credentials against DB.", "Achieved (login.php with prepared statements)"),
        ("Dashboard", "Personalized user dashboard displaying member details and track.", "Achieved (dashboard.php with profile card)"),
        ("Logout Function", "Secure logout destroying session and clearing cookies.", "Achieved (logout.php session destruction)"),
        ("Security Hashing", "Use password hashing for storing user passwords securely.", "Achieved (password_hash bcrypt standard)"),
        ("Copyright Footer", "Include required format: © First Last, Program, Sem, ID, Year.", "Achieved (footer.php exact compliance)")
    ]

    for i, row in enumerate(rubric_data):
        r_cells = rubric_table.rows[i + 1].cells
        r_cells[0].width = Inches(2.0)
        r_cells[1].width = Inches(2.6)
        r_cells[2].width = Inches(1.9)

        for j, val in enumerate(row):
            cell = r_cells[j]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = "Tahoma"
            r.font.size = Pt(9.5)
            if j == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(5, 150, 105)
            else:
                r.font.color.rgb = RGBColor(30, 41, 59)
            bg = "FFF1F2" if i % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    # Student Declaration
    p_dec = doc.add_paragraph()
    p_dec.paragraph_format.space_before = Pt(18)
    p_dec.paragraph_format.space_after = Pt(6)
    r_dec_h = p_dec.add_run("Student Declaration of Original Work:")
    r_dec_h.font.name = "Tahoma"
    r_dec_h.font.size = Pt(11)
    r_dec_h.font.bold = True
    r_dec_h.font.color.rgb = RGBColor(15, 23, 42)

    p_dec_text = doc.add_paragraph()
    p_dec_text.paragraph_format.space_after = Pt(16)
    p_dec_text.paragraph_format.line_spacing = 1.15
    r_dec_t = p_dec_text.add_run(
        "I hereby declare that this assignment titled 'Design and Development of Web Platform for Django Girls Sierra Leone' "
        "is my own original individual work produced for the Web Programming module. All code, design structures, and documentation "
        "have been prepared in accordance with academic integrity guidelines of Limkokwing University of Creative Technology Sierra Leone."
    )
    r_dec_t.font.name = "Tahoma"
    r_dec_t.font.size = Pt(10.5)
    r_dec_t.font.italic = True
    r_dec_t.font.color.rgb = RGBColor(71, 85, 105)

    # Signature Block Table
    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False

    sig_cells = sig_table.rows[0].cells
    sig_cells[0].width = Inches(3.2)
    sig_cells[1].width = Inches(3.2)
    
    p_s1 = sig_cells[0].paragraphs[0]
    p_s1.add_run("Student Signature: ______________________\nName: Alimatu Maliaka Jalloh").font.size = Pt(10)
    
    p_s2 = sig_cells[1].paragraphs[0]
    p_s2.add_run("Date: 12 October 2026\nStudent ID: 905004463").font.size = Pt(10)

    for row in sig_table.rows:
        for cell in row.cells:
            set_cell_background(cell, "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

    # Save Word document
    out_path = r"c:\Users\madi\Desktop\Web_Programming_assignment\DTECH_Web_Programming_Assignment1_Alimatu_Jalloh.docx"
    doc.save(out_path)
    print(f"Document successfully created at {out_path}")

if __name__ == "__main__":
    create_document()
